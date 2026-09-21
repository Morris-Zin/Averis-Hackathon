"""Private document spreadsheet implementation."""

from __future__ import annotations

import re
from collections.abc import Sequence
from io import BytesIO
from typing import Protocol, cast

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell

from averis.documents.evidence import block_id
from averis.documents.limits import (
    MAX_XLSX_CELLS,
    MAX_XLSX_COLUMNS,
    MAX_XLSX_ROWS_PER_SHEET,
    MAX_XLSX_SHEETS,
)
from averis.domain import DocumentEvidence, EvidenceBlock, Location


class _CalculationProperties(Protocol):
    calcMode: str | None
    fullCalcOnLoad: bool | None
    forceFullCalc: bool | None


def _spreadsheet_row(
    sheet: str,
    row: Sequence[Cell | ReadOnlyCell | EmptyCell | MergedCell],
    cached_row: Sequence[Cell | ReadOnlyCell | EmptyCell | MergedCell],
    cache_is_current: bool,
) -> tuple[list[str], list[Location], list[str]]:
    """Read a row without treating formula expressions or stale caches as data."""
    entries: list[str] = []
    locations: list[Location] = []
    uncertain: list[str] = []
    for cell, cached_cell in zip(row, cached_row, strict=True):
        if isinstance(cell, (EmptyCell, MergedCell)) or cell.value is None:
            continue
        value = str(cell.value)
        if cell.data_type == "f" or value.startswith("="):
            if not cache_is_current or not _usable_cached_formula(
                cached_cell.value, cached_cell.data_type
            ):
                uncertain.append(f"{sheet}!{cell.coordinate}")
                continue
            value = str(cached_cell.value)
        entries.append(value)
        locations.append(Location(kind="xlsx", sheet=sheet, cell=cell.coordinate))
    return entries, locations, uncertain


def read_xlsx(document_id: str, content: bytes) -> DocumentEvidence:
    from openpyxl import load_workbook

    result = DocumentEvidence(document_id=document_id)
    workbook = load_workbook(
        BytesIO(content),
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        cached_workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception:
        workbook.close()
        raise
    block_number = 0
    scanned_cells = 0
    uncertain_formula_cells: list[str] = []
    calculation = cast(_CalculationProperties, workbook.calculation)
    cache_is_current = (
        calculation.calcMode in {None, "auto"}
        and calculation.fullCalcOnLoad is not True
        and calculation.forceFullCalc is not True
    )
    try:
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            result.issues.append(
                f"xlsx_sheet_limit_exceeded:{len(workbook.worksheets)}>{MAX_XLSX_SHEETS}"
            )
        for worksheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
            cached_worksheet = cached_workbook[worksheet.title]
            title_block_added = _meaningful_worksheet_title(worksheet.title)
            if title_block_added:
                block_number += 1
                result.blocks.append(
                    EvidenceBlock(
                        id=block_id(document_id, block_number),
                        text=f"Worksheet: {worksheet.title}",
                        locations=[Location(kind="xlsx", sheet=worksheet.title)],
                    )
                )
            sheet_has_cell_evidence = False
            max_row = min(worksheet.max_row or 0, MAX_XLSX_ROWS_PER_SHEET)
            max_column = min(worksheet.max_column or 0, MAX_XLSX_COLUMNS)
            if (worksheet.max_row or 0) > MAX_XLSX_ROWS_PER_SHEET:
                result.issues.append(
                    f"xlsx_row_limit_exceeded:{worksheet.title}:{worksheet.max_row}"
                )
            if (worksheet.max_column or 0) > MAX_XLSX_COLUMNS:
                result.issues.append(
                    f"xlsx_column_limit_exceeded:{worksheet.title}:{worksheet.max_column}"
                )
            remaining_cells = MAX_XLSX_CELLS - scanned_cells
            bounded_rows = remaining_cells // max(max_column, 1)
            if max_row > bounded_rows:
                result.issues.append(
                    "xlsx_cell_scan_limit_exceeded:"
                    f"{worksheet.title}:{max_row * max_column}>{remaining_cells}"
                )
                max_row = bounded_rows
            if max_row <= 0:
                if title_block_added:
                    result.blocks.pop()
                    block_number -= 1
                break
            scanned_cells += max_row * max_column
            formula_rows = worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_column,
            )
            cached_rows = cached_worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_column,
            )
            for row, cached_row in zip(formula_rows, cached_rows, strict=True):
                entries, locations, uncertain = _spreadsheet_row(
                    worksheet.title, row, cached_row, cache_is_current
                )
                uncertain_formula_cells.extend(uncertain)
                row_has_uncertain_formula = bool(uncertain)
                # A formula expression is code, not shipment evidence. Withhold the
                # whole row when any result is stale or absent so it cannot match.
                if entries and not row_has_uncertain_formula:
                    block_number += 1
                    sheet_has_cell_evidence = True
                    result.blocks.append(
                        EvidenceBlock(
                            id=block_id(document_id, block_number),
                            text=_format_spreadsheet_row(entries),
                            locations=locations,
                        )
                    )
            if title_block_added and not sheet_has_cell_evidence:
                result.blocks.pop()
                block_number -= 1
    finally:
        workbook.close()
        cached_workbook.close()

    if uncertain_formula_cells:
        sample = ",".join(uncertain_formula_cells[:20])
        suffix = "" if len(uncertain_formula_cells) <= 20 else ",..."
        result.issues.append(
            "xlsx_formula_values_uncertain:"
            f"{len(uncertain_formula_cells)}:{sample}{suffix}"
        )
    if not result.blocks:
        result.issues.append("document_has_no_readable_text")
    return result


def _usable_cached_formula(value: object, data_type: str) -> bool:
    if value is None or data_type == "e":
        return False
    return not isinstance(value, str) or bool(value.strip())


def _meaningful_worksheet_title(title: str) -> bool:
    normalized = re.sub(r"[\s._-]+", "", title).casefold()
    return normalized not in {"sheet", "sheet1", "worksheet", "worksheet1"}


def _format_spreadsheet_row(values: Sequence[str]) -> str:
    if len(values) == 2:
        return f"{values[0]}: {values[1]}"
    return "\t".join(values)
