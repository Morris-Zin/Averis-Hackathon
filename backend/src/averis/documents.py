"""Bounded, source-linked document reading.

The public functions in this module deliberately return domain evidence rather
than objects from any of the underlying parser libraries.  A caller therefore
does not need to know whether text came from a text decoder, Office archive,
PDF layout recovery, or OCR.
"""

from __future__ import annotations

import math
import multiprocessing
import os
import re
import signal
import sys
import unicodedata
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from importlib import import_module
from io import BytesIO
from multiprocessing.connection import Connection
from pathlib import PurePath
from typing import Any, Protocol, cast
from zipfile import BadZipFile, ZipFile

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from PIL import Image as Pillow
from PIL.Image import Image as PillowImage

from averis.domain import DocumentEvidence, EvidenceBlock, Location
from averis.fields import DOCUMENT_BOUNDARY_LABELS
from averis.fields import FIELD_ALIASES as _SHARED_FIELD_ALIASES
from averis.ocr import OcrLine, read_page

MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_ARCHIVE_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_ENTRY_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 10_000
MAX_PDF_PAGES = 20
MAX_OCR_PAGES = 3
MAX_RENDER_PIXELS = 4_000_000
MAX_RENDER_DIMENSION = 3_000
MAX_XLSX_SHEETS = 50
MAX_XLSX_ROWS_PER_SHEET = 20_000
MAX_XLSX_COLUMNS = 256
MAX_XLSX_CELLS = 200_000
DEFAULT_TIMEOUT_SECONDS = 30.0

_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}
_TEXT_LABEL = re.compile(r"^[^\s:\r\n][^:\r\n]{0,80}:\s*\S")

# OCR engine selection is hidden behind the page-reading interface.
from averis.versions import OCR_PROFILE, READER_VERSION


def validate_document_evidence(evidence: DocumentEvidence) -> list[str]:
    """Validate identity, unique IDs, locations and partial-reading outcomes.

    Returns a list of completeness problems; empty means the evidence is
    structurally valid. Shipping values are never interpreted here.
    """

    problems: list[str] = []
    if not evidence.document_id:
        problems.append("missing_document_identity")
    seen: set[str] = set()
    for block in evidence.blocks:
        if block.id in seen:
            problems.append(f"duplicate_evidence_id:{block.id}")
        seen.add(block.id)
        if not block.text.strip():
            problems.append(f"empty_evidence_text:{block.id}")
        if not block.locations:
            problems.append(f"missing_location:{block.id}")
        for location in block.locations:
            if location.kind == "docx" and location.page is not None:
                problems.append(f"invented_docx_page:{block.id}")
    if not evidence.blocks and not evidence.issues:
        problems.append("empty_reading_without_issue")
    return problems


class _PdfBitmap(Protocol):
    def to_pil(self) -> PillowImage: ...

    def close(self) -> None: ...


class _PdfPage(Protocol):
    def get_size(self) -> tuple[float, float]: ...

    def render(self, *, scale: float) -> _PdfBitmap: ...

    def close(self) -> None: ...


class _PdfDocument(Protocol):
    def __len__(self) -> int: ...

    def __getitem__(self, page: int) -> _PdfPage: ...

    def close(self) -> None: ...


class _PdfiumModule(Protocol):
    PdfDocument: Callable[[bytes], _PdfDocument]


class _CalculationProperties(Protocol):
    calcMode: str | None
    fullCalcOnLoad: bool | None
    forceFullCalc: bool | None


class _ResourceModule(Protocol):
    RLIMIT_AS: int
    RLIMIT_CPU: int

    def setrlimit(self, resource: int, limits: tuple[int, int]) -> None: ...


class DocumentPreviewError(RuntimeError):
    """A stable, user-safe preview failure code."""


class _ChildProcess(Protocol):
    @property
    def pid(self) -> int | None: ...

    def is_alive(self) -> bool: ...

    def terminate(self) -> None: ...

    def kill(self) -> None: ...

    def join(self, timeout: float | None = None) -> None: ...


def _acquire_child_process_tree() -> int | None:
    """Put this worker and descendants in one OS-owned termination unit."""

    if sys.platform != "win32":
        os.setsid()
        return None

    import ctypes
    from ctypes import wintypes

    class _IoCounters(ctypes.Structure):
        _fields_ = [
            ("ReadOperationCount", ctypes.c_ulonglong),
            ("WriteOperationCount", ctypes.c_ulonglong),
            ("OtherOperationCount", ctypes.c_ulonglong),
            ("ReadTransferCount", ctypes.c_ulonglong),
            ("WriteTransferCount", ctypes.c_ulonglong),
            ("OtherTransferCount", ctypes.c_ulonglong),
        ]

    class _BasicLimitInformation(ctypes.Structure):
        _fields_ = [
            ("PerProcessUserTimeLimit", ctypes.c_longlong),
            ("PerJobUserTimeLimit", ctypes.c_longlong),
            ("LimitFlags", wintypes.DWORD),
            ("MinimumWorkingSetSize", ctypes.c_size_t),
            ("MaximumWorkingSetSize", ctypes.c_size_t),
            ("ActiveProcessLimit", wintypes.DWORD),
            ("Affinity", ctypes.c_size_t),
            ("PriorityClass", wintypes.DWORD),
            ("SchedulingClass", wintypes.DWORD),
        ]

    class _ExtendedLimitInformation(ctypes.Structure):
        _fields_ = [
            ("BasicLimitInformation", _BasicLimitInformation),
            ("IoInfo", _IoCounters),
            ("ProcessMemoryLimit", ctypes.c_size_t),
            ("JobMemoryLimit", ctypes.c_size_t),
            ("PeakProcessMemoryUsed", ctypes.c_size_t),
            ("PeakJobMemoryUsed", ctypes.c_size_t),
        ]

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_job = kernel32.CreateJobObjectW
    create_job.argtypes = [wintypes.LPVOID, wintypes.LPCWSTR]
    create_job.restype = wintypes.HANDLE
    set_information = kernel32.SetInformationJobObject
    set_information.argtypes = [
        wintypes.HANDLE,
        ctypes.c_int,
        wintypes.LPVOID,
        wintypes.DWORD,
    ]
    set_information.restype = wintypes.BOOL
    assign_process = kernel32.AssignProcessToJobObject
    assign_process.argtypes = [wintypes.HANDLE, wintypes.HANDLE]
    assign_process.restype = wintypes.BOOL
    current_process = kernel32.GetCurrentProcess
    current_process.argtypes = []
    current_process.restype = wintypes.HANDLE
    close_handle = kernel32.CloseHandle
    close_handle.argtypes = [wintypes.HANDLE]
    close_handle.restype = wintypes.BOOL

    job = create_job(None, None)
    if not job:
        raise ctypes.WinError(ctypes.get_last_error())
    information = _ExtendedLimitInformation()
    information.BasicLimitInformation.LimitFlags = 0x00002000
    if not set_information(
        job, 9, ctypes.byref(information), ctypes.sizeof(information)
    ):
        error = ctypes.WinError(ctypes.get_last_error())
        close_handle(job)
        raise error
    if not assign_process(job, current_process()):
        error = ctypes.WinError(ctypes.get_last_error())
        close_handle(job)
        raise error
    return int(job)


def _stop_process_tree(process: _ChildProcess) -> None:
    """Stop the worker and every subprocess that belongs to its ownership unit."""

    if sys.platform != "win32" and process.pid is not None:
        group_was_created = True
        try:
            os.killpg(process.pid, signal.SIGTERM)
        except ProcessLookupError:
            group_was_created = False
            # A timeout can win the race before the worker calls setsid().
            if process.is_alive():
                process.terminate()
        process.join(timeout=2)

        if group_was_created:
            # Descendants may ignore SIGTERM after the group leader has exited.
            # Always send SIGKILL to the group instead of gating it on the leader.
            try:
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
        elif process.is_alive():
            process.kill()
        process.join(timeout=2)
        return

    if process.is_alive():
        # The worker owns a kill-on-close Windows Job Object. Terminating it closes
        # that handle, and Windows terminates every process assigned to the job.
        process.terminate()
    process.join(timeout=2)
    if process.is_alive():
        process.kill()
        process.join(timeout=2)


def _with_metadata(
    document_id: str, content: bytes, evidence: DocumentEvidence
) -> DocumentEvidence:
    """Stamp immutable source identity, reader version and evidence fingerprint."""

    from hashlib import sha256

    from averis.domain import evidence_fingerprint

    evidence.document_id = document_id
    evidence.reader_version = READER_VERSION
    evidence.ocr_profile = OCR_PROFILE
    try:
        evidence.source_sha256 = sha256(content).hexdigest()
    except Exception:  # noqa: BLE001 - hashing never blocks evidence return
        evidence.source_sha256 = None
    try:
        evidence.evidence_fingerprint = evidence_fingerprint(evidence)
    except Exception:  # noqa: BLE001 - fingerprint failures stay explicit
        evidence.evidence_fingerprint = None
    return evidence


def read_document(document_id: str, filename: str, content: bytes) -> DocumentEvidence:
    """Read one supported attachment into verbatim, source-linked evidence.

    Failure and truncation are part of the result contract.  This function
    never interprets an unreadable or unsupported document as complete.
    For hostile/untrusted inputs, application code should normally call
    :func:`read_document_bounded`, which adds a process time boundary.
    """

    result = DocumentEvidence(document_id=document_id)

    if not content:
        result.issues.append("empty_document")
        return _with_metadata(document_id, content, result)
    if len(content) > MAX_DOCUMENT_BYTES:
        result.issues.append(
            f"document_size_limit_exceeded:{len(content)}>{MAX_DOCUMENT_BYTES}"
        )
        return _with_metadata(document_id, content, result)

    suffix = PurePath(filename).suffix.lower()
    adapter = FORMAT_READERS.get(suffix)
    if adapter is None:
        result.issues.append(f"unsupported_document_type:{suffix or 'none'}")
        return _with_metadata(document_id, content, result)

    try:
        if adapter.office_archive:
            archive_issue = _check_office_archive(content)
            if archive_issue:
                result.issues.append(archive_issue)
                return _with_metadata(document_id, content, result)
        return _with_metadata(
            document_id, content, adapter.read(document_id, content, suffix)
        )
    except Exception as exc:  # noqa: BLE001 - failures are result data at this boundary.
        result.issues.append(f"document_unreadable:{type(exc).__name__}")
        return _with_metadata(document_id, content, result)


def read_document_bounded(
    document_id: str,
    filename: str,
    content: bytes,
    timeout_seconds: float = DEFAULT_TIMEOUT_SECONDS,
) -> DocumentEvidence:
    """Read a document in an isolated child process with a hard wall timeout."""

    if not math.isfinite(timeout_seconds) or timeout_seconds <= 0:
        return DocumentEvidence(
            document_id=document_id,
            issues=["invalid_timeout:must_be_positive_and_finite"],
        )
    # Avoid copying an already-rejected large document into a child process.
    if len(content) > MAX_DOCUMENT_BYTES:
        return read_document(document_id, filename, content)

    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(
        target=_bounded_worker,
        args=(send, document_id, filename, content),
        daemon=True,
    )
    process.start()
    send.close()
    try:
        if receive.poll(timeout_seconds):
            try:
                payload = receive.recv_bytes()
            except (EOFError, OSError):
                return DocumentEvidence(
                    document_id=document_id,
                    issues=["document_reader_failed:no_result"],
                )
            return DocumentEvidence.model_validate_json(payload)

        return DocumentEvidence(
            document_id=document_id,
            issues=[f"document_read_timeout:{timeout_seconds:g}s"],
        )
    finally:
        receive.close()
        _stop_process_tree(process)


def render_preview_bounded(
    filename: str,
    content: bytes,
    page: int = 1,
    timeout_seconds: float = 20.0,
) -> bytes:
    """Render a PDF page or image document to a bounded PNG in a child process.

    Unsupported formats, invalid pages, corrupt documents, and timeouts raise
    :class:`DocumentPreviewError` with a stable failure code suitable for an API
    error response.  No source document content is included in the error.
    """

    if not content:
        raise DocumentPreviewError("preview_empty_document")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise DocumentPreviewError(
            f"preview_document_size_limit_exceeded:{len(content)}>{MAX_DOCUMENT_BYTES}"
        )
    suffix = PurePath(filename).suffix.lower()
    if suffix != ".pdf" and suffix not in _IMAGE_SUFFIXES:
        raise DocumentPreviewError(
            f"preview_unsupported_document_type:{suffix or 'none'}"
        )
    if page < 1 or page > MAX_PDF_PAGES:
        raise DocumentPreviewError(f"preview_page_out_of_range:{page}")
    if suffix in _IMAGE_SUFFIXES and page != 1:
        raise DocumentPreviewError(f"preview_page_out_of_range:{page}")
    if not math.isfinite(timeout_seconds) or timeout_seconds <= 0:
        raise DocumentPreviewError("preview_invalid_timeout")

    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(
        target=_preview_worker,
        args=(send, content, page, suffix),
        daemon=True,
    )
    process.start()
    send.close()
    try:
        if receive.poll(timeout_seconds):
            try:
                payload = receive.recv_bytes()
            except (EOFError, OSError) as exc:
                raise DocumentPreviewError("preview_renderer_failed:no_result") from exc
            if payload[:1] == b"\x00":
                return payload[1:]
            message = payload[1:].decode("utf-8", errors="replace")
            raise DocumentPreviewError(message or "preview_renderer_failed:no_result")

        raise DocumentPreviewError(f"preview_timeout:{timeout_seconds:g}s")
    finally:
        receive.close()
        _stop_process_tree(process)


def _apply_child_resource_limits(cpu_seconds: int) -> None:
    """Add OS-enforced child limits on POSIX; Windows retains parent wall limits."""

    if sys.platform == "win32":
        return
    resources = cast(_ResourceModule, import_module("resource"))
    # ONNX maps model/workspace memory beyond its resident set. The multilingual
    # reader is measured separately inside a 1 GiB container; retain a finite
    # address-space ceiling rather than rejecting valid model allocations.
    memory_bytes = 1536 * 1024 * 1024
    resources.setrlimit(resources.RLIMIT_AS, (memory_bytes, memory_bytes))
    resources.setrlimit(resources.RLIMIT_CPU, (cpu_seconds, cpu_seconds + 1))


def _preview_worker(
    connection: Connection,
    content: bytes,
    page: int,
    suffix: str,
) -> None:
    try:
        _acquire_child_process_tree()
    except OSError:
        connection.send_bytes(b"\x01preview_process_tree_isolation_unavailable")
        connection.close()
        return
    _apply_child_resource_limits(20)
    try:
        try:
            preview = (
                _render_pdf_preview(content, page)
                if suffix == ".pdf"
                else _render_image_preview(content, suffix)
            )
            connection.send_bytes(b"\x00" + preview)
        except Exception as exc:  # noqa: BLE001 - isolated parser failures become safe codes.
            if isinstance(exc, DocumentPreviewError):
                issue = str(exc)
            else:
                issue = f"preview_unreadable:{type(exc).__name__}"
            connection.send_bytes(b"\x01" + issue.encode("utf-8"))
    finally:
        connection.close()


def _render_pdf_preview(content: bytes, page_number: int) -> bytes:
    pdfium = cast(_PdfiumModule, import_module("pypdfium2"))

    document = pdfium.PdfDocument(content)
    try:
        if page_number > len(document):
            raise DocumentPreviewError(f"preview_page_out_of_range:{page_number}")
        page = document[page_number - 1]
        try:
            width_points, height_points = page.get_size()
            scale = min(
                150 / 72,
                MAX_RENDER_DIMENSION / max(width_points, height_points),
                math.sqrt(MAX_RENDER_PIXELS / (width_points * height_points)),
            )
            if not math.isfinite(scale) or scale <= 0:
                raise DocumentPreviewError(
                    f"preview_page_invalid_dimensions:{page_number}"
                )
            bitmap = page.render(scale=scale)  # pyright: ignore[reportArgumentType]
            try:
                image = bitmap.to_pil()
                try:
                    output = BytesIO()
                    image.save(output, format="PNG")
                    return output.getvalue()
                finally:
                    image.close()
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        document.close()


def _render_image_preview(content: bytes, suffix: str) -> bytes:
    expected_format = "PNG" if suffix == ".png" else "JPEG"
    with Pillow.open(BytesIO(content)) as source:
        if source.format != expected_format:
            raise DocumentPreviewError(f"preview_unreadable:invalid_{suffix[1:]}")
        width, height = source.size
        if (
            width <= 0
            or height <= 0
            or width > MAX_RENDER_DIMENSION
            or height > MAX_RENDER_DIMENSION
            or width * height > MAX_RENDER_PIXELS
        ):
            raise DocumentPreviewError(
                f"preview_image_pixel_limit_exceeded:{width}x{height}>"
                f"{MAX_RENDER_PIXELS}"
            )
        source.load()
        image = source.convert("RGB")
    try:
        output = BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()
    finally:
        image.close()


def _bounded_worker(
    connection: Connection,
    document_id: str,
    filename: str,
    content: bytes,
) -> None:
    try:
        _acquire_child_process_tree()
    except OSError:
        result = DocumentEvidence(
            document_id=document_id,
            issues=["document_reader_failed:process_tree_isolation_unavailable"],
        )
        connection.send_bytes(result.model_dump_json().encode("utf-8"))
        connection.close()
        return
    _apply_child_resource_limits(120)
    try:
        result = read_document(document_id, filename, content)
        connection.send_bytes(result.model_dump_json().encode("utf-8"))
    finally:
        connection.close()


def _read_text(document_id: str, content: bytes) -> DocumentEvidence:
    result = DocumentEvidence(document_id=document_id)
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        result.issues.append("text_decode_failed:utf-8")
        return result

    lines = text.splitlines()
    block_number = 0
    current: list[tuple[int, str]] = []

    def flush() -> None:
        nonlocal block_number
        if not current:
            return
        block_number += 1
        result.blocks.append(
            EvidenceBlock(
                id=_block_id(document_id, block_number),
                text="\n".join(line for _, line in current),
                locations=[
                    Location(kind="text", line_start=number, line_end=number)
                    for number, _ in current
                ],
            )
        )
        current.clear()

    for number, line in enumerate(lines, start=1):
        if line.strip():
            if current and _starts_text_field(line):
                flush()
            current.append((number, line))
        else:
            flush()
    flush()
    if not result.blocks:
        result.issues.append("document_has_no_readable_text")
    return result


def _starts_text_field(line: str) -> bool:
    """Start a source block at a labelled field, including unfamiliar fields.

    Known shipment labels may omit punctuation, while an unfamiliar label must
    start at the left margin and use a colon. Indented address continuations
    such as ``P.O. BOX: 123`` therefore stay with their owning field.
    """

    return (
        _starts_evidence_field(line)
        or _TEXT_LABEL.match(unicodedata.normalize("NFKC", line)) is not None
    )


def _read_pdf(document_id: str, content: bytes) -> DocumentEvidence:
    import pdfplumber

    result = DocumentEvidence(document_id=document_id)
    ocr_candidates: list[int] = []
    block_number = 0

    with pdfplumber.open(BytesIO(content)) as pdf:
        page_count = len(pdf.pages)
        if page_count > MAX_PDF_PAGES:
            result.issues.append(
                f"pdf_page_limit_exceeded:{page_count}>{MAX_PDF_PAGES}"
            )
        for page_number, page in enumerate(pdf.pages[:MAX_PDF_PAGES], start=1):
            words = page.extract_words(
                x_tolerance=2,
                y_tolerance=3,
                keep_blank_chars=False,
                use_text_flow=False,
                # Preserve font runs before spatial grouping. A bold form label
                # may overlap its regular-text value; merging their characters
                # first can corrupt both the label and a company name.
                extra_attrs=["fontname"],
            )
            line_groups = _pdf_lines(words)
            readable = sum(len(line[0].strip()) for line in line_groups)
            # A selectable heading does not make the scanned body readable.
            # Large page images with sparse native text need the OCR path too.
            large_image = any(
                abs(float(image["x1"]) - float(image["x0"]))
                * abs(float(image["bottom"]) - float(image["top"]))
                > float(page.width) * float(page.height) * 0.25
                for image in page.images
            )
            if readable < 4 or (large_image and readable < 200):
                ocr_candidates.append(page_number)
                continue
            for text, locations in _pdf_blocks(line_groups, page_number):
                block_number += 1
                result.blocks.append(
                    EvidenceBlock(
                        id=_block_id(document_id, block_number),
                        text=text,
                        locations=locations,
                    )
                )

    if ocr_candidates:
        if len(ocr_candidates) > MAX_OCR_PAGES:
            result.issues.append(
                f"pdf_ocr_page_limit_exceeded:{len(ocr_candidates)}>{MAX_OCR_PAGES}"
            )
        try:
            ocr_blocks, ocr_issues = _ocr_pdf_pages(
                document_id,
                content,
                ocr_candidates[:MAX_OCR_PAGES],
                block_number,
            )
            result.blocks.extend(ocr_blocks)
            result.issues.extend(ocr_issues)
        except Exception as exc:  # noqa: BLE001 - OCR failures are explicit.
            result.issues.append(f"pdf_ocr_unavailable:{type(exc).__name__}")

    if not result.blocks:
        result.issues.append("document_has_no_readable_text")
    return result


def _pdf_lines(
    words: Sequence[dict[str, Any]],
) -> list[tuple[str, tuple[float, float, float, float]]]:
    ordered = sorted(words, key=lambda word: (float(word["top"]), float(word["x0"])))
    rows: list[list[dict[str, Any]]] = []
    for word in ordered:
        if not rows or abs(float(word["top"]) - float(rows[-1][0]["top"])) > 3:
            rows.append([word])
        else:
            rows[-1].append(word)

    lines: list[tuple[str, tuple[float, float, float, float]]] = []
    for row in rows:
        row.sort(key=lambda word: float(word["x0"]))
        text = _pdf_line_text(row)
        if not text:
            continue
        lines.append(
            (
                text,
                (
                    min(float(word["x0"]) for word in row),
                    min(float(word["top"]) for word in row),
                    max(float(word["x1"]) for word in row),
                    max(float(word["bottom"]) for word in row),
                ),
            )
        )
    return lines


def _pdf_line_text(words: Sequence[dict[str, Any]]) -> str:
    """Rejoin touching font fragments, never overlapping label/value runs."""
    parts: list[str] = []
    previous: dict[str, Any] | None = None
    for word in words:
        text = str(word["text"])
        touching_font_change = (
            previous is not None
            and previous.get("fontname") != word.get("fontname")
            and -0.25 <= float(word["x0"]) - float(previous["x1"]) <= 0.5
            and str(previous["text"])[-1:].isalnum()
            and text[:1].isalnum()
        )
        parts.append(("" if touching_font_change or not parts else " ") + text)
        previous = word
    return "".join(parts).strip()


def _pdf_blocks(
    lines: Sequence[tuple[str, tuple[float, float, float, float]]],
    page_number: int,
) -> Iterable[tuple[str, list[Location]]]:
    current: list[tuple[str, tuple[float, float, float, float]]] = []
    previous_bottom: float | None = None
    for line in lines:
        top = line[1][1]
        if current and (
            _starts_evidence_field(line[0])
            or (previous_bottom is not None and top - previous_bottom > 12)
        ):
            yield _make_pdf_block(current, page_number)
            current = []
        current.append(line)
        previous_bottom = line[1][3]
    if current:
        yield _make_pdf_block(current, page_number)


def _make_pdf_block(
    lines: Sequence[tuple[str, tuple[float, float, float, float]]],
    page_number: int,
) -> tuple[str, list[Location]]:
    return (
        "\n".join(line[0] for line in lines),
        [Location(kind="pdf", page=page_number, bbox=line[1]) for line in lines],
    )


def _ocr_pdf_pages(
    document_id: str,
    content: bytes,
    page_numbers: Sequence[int],
    starting_block_number: int,
) -> tuple[list[EvidenceBlock], list[str]]:
    pdfium = cast(_PdfiumModule, import_module("pypdfium2"))

    blocks: list[EvidenceBlock] = []
    issues: list[str] = []
    block_number = starting_block_number
    document = pdfium.PdfDocument(content)
    try:
        for page_number in page_numbers:
            page = document[page_number - 1]
            try:
                width_points, height_points = page.get_size()
                scale = min(
                    200 / 72,
                    MAX_RENDER_DIMENSION / max(width_points, height_points),
                    math.sqrt(MAX_RENDER_PIXELS / (width_points * height_points)),
                )
                if not math.isfinite(scale) or scale <= 0:
                    issues.append(f"pdf_page_invalid_dimensions:{page_number}")
                    continue
                bitmap = page.render(scale=scale)  # pyright: ignore[reportArgumentType]
                try:
                    image = bitmap.to_pil()
                    try:
                        data = read_page(image)
                    finally:
                        image.close()
                finally:
                    bitmap.close()
            finally:
                page.close()

            page_blocks = _ocr_blocks(data, scale)
            if not page_blocks:
                issues.append(f"pdf_page_unreadable:{page_number}")
                continue
            for text, boxes, ocr_confidence in page_blocks:
                block_number += 1
                blocks.append(
                    EvidenceBlock(
                        id=_block_id(document_id, block_number),
                        text=text,
                        locations=[
                            Location(kind="pdf", page=page_number, bbox=box)
                            for box in boxes
                        ],
                        method="ocr",
                        ocr_confidence=ocr_confidence,
                    )
                )
    finally:
        document.close()
    return blocks, issues


def _read_image(
    document_id: str,
    content: bytes,
    suffix: str,
) -> DocumentEvidence:
    result = DocumentEvidence(document_id=document_id)
    expected_format = "PNG" if suffix == ".png" else "JPEG"
    with Pillow.open(BytesIO(content)) as source:
        if source.format != expected_format:
            result.issues.append(f"document_unreadable:invalid_{suffix[1:]}")
            return result
        width, height = source.size
        if (
            width <= 0
            or height <= 0
            or width > MAX_RENDER_DIMENSION
            or height > MAX_RENDER_DIMENSION
            or width * height > MAX_RENDER_PIXELS
        ):
            result.issues.append(
                f"image_pixel_limit_exceeded:{width}x{height}>{MAX_RENDER_PIXELS}"
            )
            return result
        source.load()
        image = source.convert("RGB")
    try:
        data = read_page(image)
    finally:
        image.close()

    for block_number, (text, boxes, ocr_confidence) in enumerate(
        _ocr_blocks(data), start=1
    ):
        result.blocks.append(
            EvidenceBlock(
                id=_block_id(document_id, block_number),
                text=text,
                locations=[Location(kind="image", bbox=box) for box in boxes],
                method="ocr",
                ocr_confidence=ocr_confidence,
            )
        )
    if not result.blocks:
        result.issues.append("document_has_no_readable_text")
    return result


def _ocr_blocks(
    recognized_lines: list[OcrLine],
    scale: float = 1.0,
) -> list[tuple[str, list[tuple[float, float, float, float]], float | None]]:
    paragraphs: dict[
        tuple[int, int],
        list[tuple[int, str, tuple[float, float, float, float], float | None]],
    ] = {}
    for line in recognized_lines:
        left, top, right, bottom = line.bbox
        paragraphs.setdefault(line.paragraph, []).append(
            (
                line.order,
                line.text,
                (left / scale, top / scale, right / scale, bottom / scale),
                line.confidence,
            )
        )

    blocks: list[tuple[str, list[tuple[float, float, float, float]], float | None]] = []
    for lines in paragraphs.values():
        lines.sort(key=lambda line: (line[0], line[2][1], line[2][0]))
        current: list[
            tuple[int, str, tuple[float, float, float, float], float | None]
        ] = []
        for line in lines:
            if current and _starts_evidence_field(line[1]):
                blocks.append(_make_ocr_block(current))
                current = []
            current.append(line)
        if current:
            blocks.append(_make_ocr_block(current))
    return blocks


_EVIDENCE_FIELD_LABEL = re.compile(
    r"^\s*(?:"
    r"shipper|exporter|consignee|notify\s*party|notify|"
    r"port\s*of\s*loading|load\s+port|pol|"
    r"port\s*of\s*discharge|discharge\s+port|pod|"
    r"number\s+of\s+containers(?:\s+or\s+packages)?|"
    r"no\.\s*of\s+containers(?:\s+or\s+packages)?|"
    r"container\s+count|containers|"
    r"(?:total\s+)?gross\s+weight(?:毛重)?(?:\s*\([^)]*\)|\s+kgs?)?|"
    r"(?:total\s+)?gross\s+wt(?:\s*\([^)]*\)|\s+kgs?)?|"
    r"b/l\s+(?:no|number)|bill\s+of\s+lading\s+(?:no|number)|"
    r"booking(?:\s+(?:no|number|ref|reference))?|"
    r"ocean\s+vessel|vessel|voyage|export\s+carrier|container\s+no|hs\s+code"
    r")(?=\s*[:=|]\s*\S|\.?\s+\S)",
    re.IGNORECASE,
)


def official_field_aliases() -> dict[str, tuple[str, ...]]:
    """Return the shared official aliases; layout grouping stays private.

    The segmentation regex above covers these official aliases plus
    parser-only headings (vessel, voyage, booking, B/L number, HS code).
    Multiline addresses are preserved because only a new label starts a block.
    """

    return {str(key): value for key, value in _SHARED_FIELD_ALIASES.items()}


def _starts_evidence_field(text: str) -> bool:
    """Recognize a new official field without dropping continuation lines."""

    normalized = unicodedata.normalize("NFKC", text)
    return (
        _EVIDENCE_FIELD_LABEL.match(normalized) is not None
        or _SHARED_FIELD_LABEL.match(normalized) is not None
        or _DOCUMENT_BOUNDARY_LABEL.match(normalized) is not None
    )


_DOCUMENT_BOUNDARY_LABEL = re.compile(
    r"^\s*(?:"
    + "|".join(
        re.escape(label)
        for label in sorted(DOCUMENT_BOUNDARY_LABELS, key=len, reverse=True)
    )
    + r")(?=\s|[:=|]|$)",
    re.IGNORECASE,
)


_SHARED_FIELD_LABEL = re.compile(
    r"^\s*(?:"
    + "|".join(
        re.escape(label)
        for label in sorted(
            {label for aliases in _SHARED_FIELD_ALIASES.values() for label in aliases},
            key=len,
            reverse=True,
        )
    )
    + r")(?:\s*\([^()\r\n]*\))*(?=\s*[:=|]\s*\S|\s+\S)",
    re.IGNORECASE,
)


def _make_ocr_block(
    lines: Sequence[tuple[int, str, tuple[float, float, float, float], float | None]],
) -> tuple[str, list[tuple[float, float, float, float]], float | None]:
    confidences = [line[3] for line in lines]
    confidence = (
        min(cast(list[float], confidences))
        if confidences and all(value is not None for value in confidences)
        else None
    )
    return (
        "\n".join(line[1] for line in lines),
        [line[2] for line in lines],
        confidence,
    )


def _read_docx(document_id: str, content: bytes) -> DocumentEvidence:
    from averis.word_structure import read_word_content

    result = DocumentEvidence(document_id=document_id)
    document = read_word_content(content)
    result.issues.extend(document.issues)
    block_number = 0
    field_paragraphs: list[EvidenceBlock] = []

    def flush_field() -> None:
        if not field_paragraphs:
            return
        first = field_paragraphs[0]
        result.blocks.append(
            first.model_copy(
                update={
                    "text": "\n".join(block.text for block in field_paragraphs),
                    "locations": [
                        location
                        for block in field_paragraphs
                        for location in block.locations
                    ],
                }
            )
        )
        field_paragraphs.clear()

    for item in document.items:
        if item.boundary:
            flush_field()
            continue
        if item.table:
            flush_field()
            cell_texts: list[str] = []
            locations: list[Location] = []
            for cell in item.cells:
                paragraphs: list[str] = []
                for paragraph in cell:
                    if paragraph.text.strip():
                        paragraphs.append(paragraph.text)
                        locations.append(
                            Location(kind="docx", paragraph=paragraph.number)
                        )
                if paragraphs:
                    cell_texts.append("\n".join(paragraphs))
            if cell_texts:
                block_number += 1
                result.blocks.append(
                    EvidenceBlock(
                        id=_block_id(document_id, block_number),
                        text=" | ".join(cell_texts),
                        locations=locations,
                    )
                )
        else:
            paragraph = item.cells[0][0]
            if paragraph.text.strip():
                block_number += 1
                block = EvidenceBlock(
                    id=_block_id(document_id, block_number),
                    text=paragraph.text,
                    locations=[Location(kind="docx", paragraph=paragraph.number)],
                )
                label_only = any(
                    re.fullmatch(
                        r"\s*" + re.escape(label) + r"\s*:?[ \t]*",
                        paragraph.text,
                        re.IGNORECASE,
                    )
                    for aliases in _SHARED_FIELD_ALIASES.values()
                    for label in aliases
                )
                if _starts_text_field(paragraph.text) or label_only:
                    flush_field()
                party_label = any(
                    re.match(
                        r"^\s*" + re.escape(label) + r"(?=\s|[:=|]|$)",
                        unicodedata.normalize("NFKC", paragraph.text),
                        re.IGNORECASE,
                    )
                    for field in ("shipper", "consignee", "notify_party")
                    for label in _SHARED_FIELD_ALIASES[field]
                )
                ambiguous_label_only = (
                    re.fullmatch(
                        r"\s*container\s+(?:number|no\.?|id)\s*[:=|]?\s*",
                        unicodedata.normalize("NFKC", paragraph.text),
                        re.IGNORECASE,
                    )
                    is not None
                )
                if party_label or ambiguous_label_only or field_paragraphs:
                    field_paragraphs.append(block)
                else:
                    result.blocks.append(block)

    flush_field()

    if not result.blocks:
        result.issues.append("document_has_no_readable_text")
    return result


def _spreadsheet_row(
    sheet: str,
    row: Sequence[Cell | ReadOnlyCell | EmptyCell | MergedCell],
    cached_row: Sequence[Cell | ReadOnlyCell | EmptyCell | MergedCell],
    cache_is_current: bool,
) -> tuple[list[str], list[Location], list[str]]:
    """Read a row without treating formula expressions or stale caches as data."""
    entries: list[str] = []
    locations: list[Location] = []
    uncertain: list[str] = []
    for cell, cached_cell in zip(row, cached_row, strict=True):
        if isinstance(cell, (EmptyCell, MergedCell)) or cell.value is None:
            continue
        value = str(cell.value)
        if cell.data_type == "f" or value.startswith("="):
            if not cache_is_current or not _usable_cached_formula(
                cached_cell.value, cached_cell.data_type
            ):
                uncertain.append(f"{sheet}!{cell.coordinate}")
                continue
            value = str(cached_cell.value)
        entries.append(value)
        locations.append(Location(kind="xlsx", sheet=sheet, cell=cell.coordinate))
    return entries, locations, uncertain


def _read_xlsx(document_id: str, content: bytes) -> DocumentEvidence:
    from openpyxl import load_workbook

    result = DocumentEvidence(document_id=document_id)
    workbook = load_workbook(
        BytesIO(content),
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        cached_workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception:
        workbook.close()
        raise
    block_number = 0
    scanned_cells = 0
    uncertain_formula_cells: list[str] = []
    calculation = cast(_CalculationProperties, workbook.calculation)
    cache_is_current = (
        calculation.calcMode in {None, "auto"}
        and calculation.fullCalcOnLoad is not True
        and calculation.forceFullCalc is not True
    )
    try:
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            result.issues.append(
                f"xlsx_sheet_limit_exceeded:{len(workbook.worksheets)}>{MAX_XLSX_SHEETS}"
            )
        for worksheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
            cached_worksheet = cached_workbook[worksheet.title]
            title_block_added = _meaningful_worksheet_title(worksheet.title)
            if title_block_added:
                block_number += 1
                result.blocks.append(
                    EvidenceBlock(
                        id=_block_id(document_id, block_number),
                        text=f"Worksheet: {worksheet.title}",
                        locations=[Location(kind="xlsx", sheet=worksheet.title)],
                    )
                )
            sheet_has_cell_evidence = False
            max_row = min(worksheet.max_row or 0, MAX_XLSX_ROWS_PER_SHEET)
            max_column = min(worksheet.max_column or 0, MAX_XLSX_COLUMNS)
            if (worksheet.max_row or 0) > MAX_XLSX_ROWS_PER_SHEET:
                result.issues.append(
                    f"xlsx_row_limit_exceeded:{worksheet.title}:{worksheet.max_row}"
                )
            if (worksheet.max_column or 0) > MAX_XLSX_COLUMNS:
                result.issues.append(
                    f"xlsx_column_limit_exceeded:{worksheet.title}:{worksheet.max_column}"
                )
            remaining_cells = MAX_XLSX_CELLS - scanned_cells
            bounded_rows = remaining_cells // max(max_column, 1)
            if max_row > bounded_rows:
                result.issues.append(
                    "xlsx_cell_scan_limit_exceeded:"
                    f"{worksheet.title}:{max_row * max_column}>{remaining_cells}"
                )
                max_row = bounded_rows
            if max_row <= 0:
                if title_block_added:
                    result.blocks.pop()
                    block_number -= 1
                break
            scanned_cells += max_row * max_column
            formula_rows = worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_column,
            )
            cached_rows = cached_worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_column,
            )
            for row, cached_row in zip(formula_rows, cached_rows, strict=True):
                entries, locations, uncertain = _spreadsheet_row(
                    worksheet.title, row, cached_row, cache_is_current
                )
                uncertain_formula_cells.extend(uncertain)
                row_has_uncertain_formula = bool(uncertain)
                # A formula expression is code, not shipment evidence. Withhold the
                # whole row when any result is stale or absent so it cannot match.
                if entries and not row_has_uncertain_formula:
                    block_number += 1
                    sheet_has_cell_evidence = True
                    result.blocks.append(
                        EvidenceBlock(
                            id=_block_id(document_id, block_number),
                            text=_format_spreadsheet_row(entries),
                            locations=locations,
                        )
                    )
            if title_block_added and not sheet_has_cell_evidence:
                result.blocks.pop()
                block_number -= 1
    finally:
        workbook.close()
        cached_workbook.close()

    if uncertain_formula_cells:
        sample = ",".join(uncertain_formula_cells[:20])
        suffix = "" if len(uncertain_formula_cells) <= 20 else ",..."
        result.issues.append(
            "xlsx_formula_values_uncertain:"
            f"{len(uncertain_formula_cells)}:{sample}{suffix}"
        )
    if not result.blocks:
        result.issues.append("document_has_no_readable_text")
    return result


def _usable_cached_formula(value: object, data_type: str) -> bool:
    if value is None or data_type == "e":
        return False
    return not isinstance(value, str) or bool(value.strip())


def _meaningful_worksheet_title(title: str) -> bool:
    normalized = re.sub(r"[\s._-]+", "", title).casefold()
    return normalized not in {"sheet", "sheet1", "worksheet", "worksheet1"}


def _format_spreadsheet_row(values: Sequence[str]) -> str:
    if len(values) == 2:
        return f"{values[0]}: {values[1]}"
    return "\t".join(values)


def _check_office_archive(content: bytes) -> str | None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                return (
                    f"archive_entry_limit_exceeded:{len(entries)}>{MAX_ARCHIVE_ENTRIES}"
                )
            expanded = 0
            for entry in entries:
                if entry.file_size > MAX_ARCHIVE_ENTRY_BYTES:
                    return (
                        "archive_member_size_limit_exceeded:"
                        f"{entry.file_size}>{MAX_ARCHIVE_ENTRY_BYTES}"
                    )
                expanded += entry.file_size
                if expanded > MAX_ARCHIVE_BYTES:
                    return f"archive_size_limit_exceeded:{expanded}>{MAX_ARCHIVE_BYTES}"
    except BadZipFile:
        return "document_unreadable:invalid_office_archive"
    return None


def _block_id(document_id: str, number: int) -> str:
    return f"{document_id}:b{number:04d}"


@dataclass(frozen=True)
class FormatReader:
    """A trusted built-in format adapter behind the bounded document boundary."""

    read: Callable[[str, bytes, str], DocumentEvidence]
    preview: bool = False
    ocr: bool = False
    office_archive: bool = False


# Static registration is imported in spawned parser processes too. Adding a
# format does not change pipeline or dispatcher logic. Keep metadata and handler
# together; uploads must explicitly authorize newly supported file extensions.
FORMAT_READERS = {
    ".txt": FormatReader(lambda doc, content, suffix: _read_text(doc, content)),
    ".pdf": FormatReader(
        lambda doc, content, suffix: _read_pdf(doc, content), True, True
    ),
    ".docx": FormatReader(
        lambda doc, content, suffix: _read_docx(doc, content), office_archive=True
    ),
    ".xlsx": FormatReader(
        lambda doc, content, suffix: _read_xlsx(doc, content), office_archive=True
    ),
    **{suffix: FormatReader(_read_image, True, True) for suffix in _IMAGE_SUFFIXES},
}
SUPPORTED_FORMATS: dict[str, dict[str, object]] = {
    suffix: {"preview": reader.preview, "ocr": reader.ocr}
    for suffix, reader in FORMAT_READERS.items()
}
