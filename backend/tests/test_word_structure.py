"""Reader regressions with independently specified Word source structures."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from averis.documents import read_document
from averis.documents.word_structure import MC, WPS, W
from averis.verification import normalize, reading_from_evidence


def package(body: str, extra: dict[str, str] | None = None) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "word/document.xml",
            f'<w:document xmlns:w="{W[1:-1]}" xmlns:mc="{MC[1:-1]}" xmlns:wps="{WPS}">'
            f"<w:body>{body}</w:body></w:document>",
        )
        for name, xml in (extra or {}).items():
            archive.writestr(name, xml)
    return buffer.getvalue()


def paragraph(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def test_textbox_compatibility_copies_are_read_once_with_complete_address() -> None:
    box = (
        "<w:txbxContent>"
        + paragraph("Draft Bill of Lading")
        + paragraph("Shipper: North Export")
        + paragraph("81 Harbour Road")
        + "<w:p><w:r><w:t>Consign</w:t></w:r>"
        "<w:r><w:t>ee: South Import</w:t></w:r></w:p></w:txbxContent>"
    )
    source = package(
        f"<w:p><w:r><mc:AlternateContent><mc:Choice Requires='wps'>{box}</mc:Choice>"
        f"<mc:Fallback>{box}</mc:Fallback></mc:AlternateContent></w:r></w:p>"
        + paragraph("Unrelated body text")
    )
    evidence = read_document("box", "bl.docx", source)
    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == [
        "Draft Bill of Lading",
        "Shipper: North Export\n81 Harbour Road",
        "Consignee: South Import",
        "Unrelated body text",
    ]
    assert [location.paragraph for location in evidence.blocks[1].locations] == [3, 4]
    assert all(
        location.page is None
        for block in evidence.blocks
        for location in block.locations
    )
    reading = reading_from_evidence("shipper", evidence, [evidence.blocks[1].id])
    assert reading.issue is None
    assert reading.normalized == "north export 81 harbour road"
    assert reading.normalized != normalize(
        "shipper", "Shipper: North Export\n80 Harbour Road"
    )


def test_content_controls_and_hyperlinks_keep_visible_text() -> None:
    source = package(
        "<w:sdt><w:sdtPr/><w:sdtContent><w:p>"
        "<w:hyperlink><w:r><w:t>Shipper: North Export</w:t></w:r></w:hyperlink>"
        "</w:p></w:sdtContent></w:sdt>"
    )
    evidence = read_document("wrapped", "si.docx", source)
    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == ["Shipper: North Export"]


def test_merged_table_cell_does_not_duplicate_evidence() -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Loading Port: Klang"
    buffer = BytesIO()
    doc.save(buffer)
    evidence = read_document("merged", "si.docx", buffer.getvalue())
    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == ["Loading Port: Klang"]
    assert normalize("port_of_loading", evidence.blocks[0].text) == "klang"


@pytest.mark.parametrize(
    ("xml", "issue"),
    [
        (
            "<w:ins>" + paragraph("Loading Port: Klang") + "</w:ins>",
            "docx_tracked_changes_require_review",
        ),
        (
            "<w:p><w:r><w:pict><w:imagedata/></w:pict></w:r></w:p>",
            "docx_embedded_content_requires_review",
        ),
        ("<w:altChunk/>", "docx_embedded_content_requires_review"),
        (
            "<w:unknown>" + paragraph("Lost text") + "</w:unknown>",
            "docx_unsupported_structure_requires_review",
        ),
        (
            "<w:p><w:r><w:rPr><w:vanish/></w:rPr><w:t>Hidden text</w:t></w:r></w:p>",
            "docx_hidden_text_requires_review",
        ),
    ],
)
def test_unhandled_visible_content_is_not_silently_complete(
    xml: str, issue: str
) -> None:
    evidence = read_document("unsupported", "si.docx", package(xml))
    assert issue in evidence.issues


def test_header_text_is_explicitly_outside_supported_preview() -> None:
    header = (
        f'<w:hdr xmlns:w="{W[1:-1]}">'
        + paragraph("Shipper: Header Export")
        + "</w:hdr>"
    )
    evidence = read_document(
        "header",
        "si.docx",
        package(paragraph("Loading Port: Klang"), {"word/header1.xml": header}),
    )
    assert "docx_external_text_requires_review" in evidence.issues


def test_nested_table_is_recovered_but_uncertain_layout_requires_review() -> None:
    nested = (
        "<w:tbl><w:tr><w:tc>"
        + paragraph("Container Count: 3")
        + "</w:tc></w:tr></w:tbl>"
    )
    body = (
        "<w:tbl><w:tr><w:tc>"
        + paragraph("Outer heading")
        + nested
        + "</w:tc></w:tr></w:tbl>"
    )
    evidence = read_document("nested", "si.docx", package(body))
    assert [block.text for block in evidence.blocks] == [
        "Outer heading",
        "Container Count: 3",
    ]
    assert "docx_nested_table_requires_review" in evidence.issues


def test_deeply_nested_textboxes_fail_with_a_bounded_reader_issue() -> None:
    body = paragraph("Container Count: 3")
    for _ in range(70):
        body = "<w:p><w:r><w:txbxContent>" + body + "</w:txbxContent></w:r></w:p>"
    evidence = read_document("deep", "si.docx", package(body))
    assert "document_unreadable:ValueError" in evidence.issues


@pytest.mark.parametrize("value", ["3", "MSCU1234567", "1234567"])
def test_container_number_does_not_silently_become_count(value: str) -> None:
    assert normalize("container_count", f"Container Number: {value}") is None


def test_loading_port_synonym_preserves_field_boundaries() -> None:
    evidence = read_document(
        "labels",
        "si.txt",
        b"Shipper: Export Co\nLoading Port: Klang\nDischarge Port: Singapore",
    )
    assert [block.text for block in evidence.blocks] == [
        "Shipper: Export Co",
        "Loading Port: Klang",
        "Discharge Port: Singapore",
    ]
    reading = reading_from_evidence(
        "port_of_loading", evidence, [evidence.blocks[1].id]
    )
    assert reading.issue is None
    assert reading.normalized == "klang"


@pytest.mark.parametrize(
    "text",
    ["Port Of Shame: Port Klang", "Port OfLame: Singapore", "Unknown label：新加坡"],
)
def test_unknown_port_label_remains_unresolved(text: str) -> None:
    assert normalize("port_of_loading", text) is None
    assert normalize("port_of_discharge", text) is None


def test_utf16_entity_declarations_cannot_expand_into_source_evidence() -> None:
    buffer = BytesIO()
    xml = (
        '<?xml version="1.0" encoding="UTF-16"?>'
        '<!DOCTYPE document [<!ENTITY invented "Invented Export">]>'
        f'<w:document xmlns:w="{W[1:-1]}"><w:body>'
        "<w:p><w:r><w:t>&invented;</w:t></w:r></w:p>"
        "</w:body></w:document>"
    )
    with ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", xml.encode("utf-16"))
    evidence = read_document("entity", "si.docx", buffer.getvalue())
    assert evidence.blocks == []
    assert "document_unreadable:ValueError" in evidence.issues


def test_two_fields_in_one_table_row_are_not_accepted_as_one_party() -> None:
    cells = "".join(
        "<w:tc>" + paragraph(text) + "</w:tc>"
        for text in ["Shipper", "ALPHA", "Consignee", "BETA"]
    )
    evidence = read_document(
        "mixed", "si.docx", package("<w:tbl><w:tr>" + cells + "</w:tr></w:tbl>")
    )
    reading = reading_from_evidence("shipper", evidence, [evidence.blocks[0].id])
    assert reading.issue == "ambiguous_source_fields"
    assert reading.normalized is None


def test_inherited_hidden_style_cannot_supply_an_unflagged_value() -> None:
    styles = (
        f'<w:styles xmlns:w="{W[1:-1]}">'
        '<w:style w:styleId="Hidden"><w:rPr><w:vanish/></w:rPr></w:style>'
        '<w:style w:styleId="Derived"><w:basedOn w:val="Hidden"/></w:style>'
        "</w:styles>"
    )
    body = (
        '<w:p><w:pPr><w:pStyle w:val="Derived"/></w:pPr>'
        "<w:r><w:t>Shipper: Invisible Export</w:t></w:r></w:p>"
    )
    evidence = read_document(
        "styled", "si.docx", package(body, {"word/styles.xml": styles})
    )
    assert "docx_hidden_text_requires_review" in evidence.issues
    unused = read_document(
        "unused",
        "si.docx",
        package(paragraph("Shipper: Visible Export"), {"word/styles.xml": styles}),
    )
    assert unused.issues == []


def test_tracked_table_row_deletion_is_not_silently_current_content() -> None:
    body = (
        "<w:tbl><w:tr><w:trPr><w:del/></w:trPr><w:tc>"
        + paragraph("Shipper: Deleted Export")
        + "</w:tc></w:tr></w:tbl>"
    )
    evidence = read_document("deleted-row", "si.docx", package(body))
    assert "docx_tracked_changes_require_review" in evidence.issues


def test_different_compatibility_representations_cannot_silently_choose_a_value() -> (
    None
):
    first = "<w:txbxContent>" + paragraph("Shipper: FIRST") + "</w:txbxContent>"
    second = "<w:txbxContent>" + paragraph("Shipper: SECOND") + "</w:txbxContent>"
    body = f"<w:p><w:r><mc:AlternateContent><mc:Choice Requires='wps'>{first}</mc:Choice><mc:Fallback>{second}</mc:Fallback></mc:AlternateContent></w:r></w:p>"
    evidence = read_document("conflict", "si.docx", package(body))
    assert len(evidence.blocks) == 1
    assert "docx_alternate_content_conflict_requires_review" in evidence.issues


@pytest.mark.parametrize(
    "label", ["Shipper = North Export", "Shipper North Export", "发货人：North Export"]
)
def test_party_addresses_do_not_depend_on_ascii_colons(label: str) -> None:
    source = package(
        paragraph(label)
        + paragraph("81 Harbour Road")
        + paragraph("Container Count: 3")
    )
    evidence = read_document("party", "si.docx", source)
    assert evidence.blocks[0].text == label + "\n81 Harbour Road"
    reading = reading_from_evidence("shipper", evidence, [evidence.blocks[0].id])
    assert reading.normalized == "north export 81 harbour road"


def test_split_ambiguous_container_label_keeps_its_value_context() -> None:
    evidence = read_document(
        "count", "si.docx", package(paragraph("Container Number:") + paragraph("3"))
    )
    assert [block.text for block in evidence.blocks] == ["Container Number:\n3"]
    assert (
        reading_from_evidence(
            "container_count", evidence, [evidence.blocks[0].id]
        ).normalized
        is None
    )


@pytest.mark.parametrize(
    "inline", ["<w:noBreakHyphen/>", '<w:sym w:font="Symbol" w:char="F02D"/>']
)
def test_displayed_characters_cannot_disappear_into_invented_weight(
    inline: str,
) -> None:
    body = (
        "<w:p><w:r><w:t>Gross weight: 10</w:t>"
        + inline
        + "<w:t>20 kg</w:t></w:r></w:p>"
    )
    evidence = read_document("weight", "si.docx", package(body))
    assert normalize("gross_weight_kg", evidence.blocks[0].text) is None


def test_compatibility_requires_resolves_namespace_uri_not_prefix() -> None:
    first = "<w:txbxContent>" + paragraph("Shipper: UNSUPPORTED") + "</w:txbxContent>"
    second = "<w:txbxContent>" + paragraph("Shipper: FALLBACK") + "</w:txbxContent>"
    body = f"<w:p><w:r><mc:AlternateContent><mc:Choice xmlns:new='urn:unsupported' Requires='new'>{first}</mc:Choice><mc:Fallback>{second}</mc:Fallback></mc:AlternateContent></w:r></w:p>"
    evidence = read_document("fallback", "si.docx", package(body))
    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == ["Shipper: FALLBACK"]
    renamed = body.replace("urn:unsupported", WPS).replace("UNSUPPORTED", "FALLBACK")
    supported = read_document("supported", "si.docx", package(renamed))
    assert supported.issues == []
    assert [block.text for block in supported.blocks] == ["Shipper: FALLBACK"]


def test_vertical_merges_require_review_instead_of_company_only_matches() -> None:
    body = (
        "<w:tbl><w:tr><w:tc><w:tcPr><w:vMerge w:val='restart'/></w:tcPr>"
        + paragraph("Shipper")
        + "</w:tc><w:tc>"
        + paragraph("North Export")
        + "</w:tc></w:tr><w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr></w:tc><w:tc>"
        + paragraph("81 Harbour Road")
        + "</w:tc></w:tr></w:tbl>"
    )
    evidence = read_document("merge", "si.docx", package(body))
    assert "docx_vertical_merge_requires_review" in evidence.issues


def test_ambiguous_container_package_count_requires_container_semantics() -> None:
    assert normalize("container_count", "Number of containers or packages: 3") is None
    assert (
        normalize("container_count", "Number of containers or packages: 3 packages")
        is None
    )
    assert (
        normalize("container_count", "Number of containers or packages: 3 containers")
        == "3"
    )
    assert (
        normalize("container_count", "Number of containers or packages: 3 x 40HQ")
        == "3"
    )


@pytest.mark.parametrize("separator", ["|", "="])
def test_unknown_port_labels_do_not_become_valid_in_other_layouts(
    separator: str,
) -> None:
    assert normalize("port_of_loading", "Port Of Shame " + separator + " Klang") is None


def test_all_required_namespaces_must_be_supported_despite_source_marker() -> None:
    body = (
        "<mc:AlternateContent><mc:Choice xmlns:new='urn:unsupported' "
        "xmlns:a='urn:averis:reader' Requires='wps new' a:supported-choice='true'>"
        + paragraph("Shipper: WRONG")
        + "</mc:Choice><mc:Fallback>"
        + paragraph("Shipper: RIGHT")
        + "</mc:Fallback></mc:AlternateContent>"
    )
    evidence = read_document("namespaces", "si.docx", package(body))
    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == ["Shipper: RIGHT"]
