import multiprocessing
import os
import subprocess
import sys
import time
from io import BytesIO
from multiprocessing.connection import Connection
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytesseract
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pytest import MonkeyPatch

from averis.documents import (
    MAX_DOCUMENT_BYTES,
    MAX_RENDER_DIMENSION,
    MAX_RENDER_PIXELS,
    _acquire_child_process_tree,
    _stop_process_tree,
    read_document,
    read_document_bounded,
    render_preview_bounded,
)
from averis.verification import normalize


def test_txt_preserves_multiline_blocks_and_line_locations() -> None:
    evidence = read_document(
        "txt-1",
        "shipment.TXT",
        "Shipper: Café Export\r\nAddress line 2\r\n\r\nWeight: 50 kg\r\n".encode(),
    )

    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == [
        "Shipper: Café Export\nAddress line 2",
        "Weight: 50 kg",
    ]
    assert [location.line_start for location in evidence.blocks[0].locations] == [1, 2]
    assert evidence.blocks[0].id == "txt-1:b0001"


def test_size_and_unsupported_formats_are_explicit_failures() -> None:
    too_large = read_document("large", "large.txt", b"x" * (MAX_DOCUMENT_BYTES + 1))
    unsupported = read_document("rtf", "shipment.rtf", b"not really rich text")

    assert too_large.blocks == []
    assert too_large.issues == [
        f"document_size_limit_exceeded:{MAX_DOCUMENT_BYTES + 1}>{MAX_DOCUMENT_BYTES}"
    ]
    assert unsupported.blocks == []
    assert unsupported.issues == ["unsupported_document_type:.rtf"]


def test_bounded_reader_returns_the_same_typed_evidence() -> None:
    evidence = read_document_bounded(
        "bounded-1",
        "bounded.txt",
        b"Shipper: Averis Trading\nAddress: Port Klang",
        timeout_seconds=10,
    )

    assert evidence.issues == []
    assert evidence.document_id == "bounded-1"
    assert evidence.blocks[0].text == "Shipper: Averis Trading\nAddress: Port Klang"


def test_bounded_pdf_preview_is_a_pixel_bounded_png() -> None:
    preview = render_preview_bounded(
        "si.pdf", _minimal_text_pdf(), page=1, timeout_seconds=10
    )

    assert preview.startswith(b"\x89PNG\r\n\x1a\n")
    with Image.open(BytesIO(preview)) as image:
        assert image.width * image.height <= MAX_RENDER_PIXELS
        assert max(image.width, image.height) <= MAX_RENDER_DIMENSION


def test_docx_keeps_paragraph_and_table_order_without_inventing_pages() -> None:
    document = Document()
    document.add_paragraph("Shipping Instruction")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Shipper"
    table.cell(0, 1).text = "Averis Trading"
    document.add_paragraph("Consignee address\nSecond line")
    payload = BytesIO()
    document.save(payload)

    evidence = read_document("docx-1", "si.docx", payload.getvalue())

    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == [
        "Shipping Instruction",
        "Shipper | Averis Trading",
        "Consignee address\nSecond line",
    ]
    assert len(evidence.blocks[1].locations) == 2
    assert all(
        location.kind == "docx"
        for block in evidence.blocks
        for location in block.locations
    )
    assert all(
        location.page is None
        for block in evidence.blocks
        for location in block.locations
    )


def test_invalid_office_archive_is_not_reported_as_empty_success() -> None:
    evidence = read_document("bad-docx", "bad.docx", b"this is not a zip archive")

    assert evidence.blocks == []
    assert evidence.issues == ["document_unreadable:invalid_office_archive"]


def test_xlsx_keeps_cell_locations_and_marks_formula_values_uncertain() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "SI"
    sheet.append(["Field", "Value"])
    sheet.append(["Gross Weight", 1250])
    sheet.append(["Calculated", "=B2*2"])
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()

    evidence = read_document("xlsx-1", "si.xlsx", payload.getvalue())

    assert len(evidence.blocks) == 2
    assert evidence.blocks[1].text == "Gross Weight: 1250"
    assert [
        (location.sheet, location.cell) for location in evidence.blocks[1].locations
    ] == [
        ("SI", "A2"),
        ("SI", "B2"),
    ]
    assert all("=B2*2" not in block.text for block in evidence.blocks)
    assert evidence.issues == ["xlsx_formula_values_uncertain:1:SI!B3"]


def test_pdf_native_text_has_page_and_bounding_box_locations() -> None:
    evidence = read_document("pdf-1", "si.pdf", _minimal_text_pdf())

    assert evidence.issues == []
    assert "Shipping Instruction" in "\n".join(block.text for block in evidence.blocks)
    locations = [location for block in evidence.blocks for location in block.locations]
    assert locations
    assert all(location.kind == "pdf" and location.page == 1 for location in locations)
    assert all(location.bbox is not None for location in locations)


def test_standalone_png_uses_bounded_ocr_with_image_location(
    monkeypatch: MonkeyPatch,
) -> None:
    image = Image.new("RGB", (300, 150), color="white")
    payload = BytesIO()
    image.save(payload, format="PNG")
    image.close()

    def fake_image_to_data(*_: object, **__: object) -> dict[str, list[object]]:
        return {
            "text": [
                "Notify",
                "Party",
                "77",
                "Robinson",
                "Gross",
                "Weight",
                "0",
                "KG",
            ],
            "block_num": [1] * 8,
            "par_num": [1] * 8,
            "line_num": [1, 1, 2, 2, 3, 3, 3, 3],
            "left": [10, 70, 10, 40, 10, 60, 120, 135],
            "top": [20, 20, 45, 45, 70, 70, 70, 70],
            "width": [50, 40, 20, 65, 45, 55, 10, 20],
            "height": [15] * 8,
        }

    monkeypatch.setattr(pytesseract, "image_to_data", fake_image_to_data)
    evidence = read_document("png-1", "scan.png", payload.getvalue())

    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == [
        "Notify Party\n77 Robinson",
        "Gross Weight 0 KG",
    ]
    assert all(block.method == "ocr" for block in evidence.blocks)
    assert evidence.blocks[0].locations[0].kind == "image"
    assert [location.bbox for location in evidence.blocks[0].locations] == [
        (10.0, 20.0, 110.0, 35.0),
        (10.0, 45.0, 105.0, 60.0),
    ]
    assert [location.bbox for location in evidence.blocks[1].locations] == [
        (10.0, 70.0, 155.0, 85.0)
    ]


def test_scanned_pdf_uses_ocr_with_page_location(monkeypatch: MonkeyPatch) -> None:
    image = Image.new("RGB", (300, 150), color="white")
    payload = BytesIO()
    image.save(payload, format="PDF")
    image.close()

    def fake_image_to_data(*_: object, **__: object) -> dict[str, list[object]]:
        return {
            "text": [
                "Shipper",
                "Averis",
                "77",
                "Robinson",
                "Consignee",
                "Meridian",
            ],
            "block_num": [1] * 6,
            "par_num": [1] * 6,
            "line_num": [1, 1, 2, 2, 3, 3],
            "left": [10, 80, 10, 40, 10, 90],
            "top": [20, 20, 45, 45, 70, 70],
            "width": [60, 50, 20, 65, 75, 70],
            "height": [15] * 6,
        }

    monkeypatch.setattr(pytesseract, "image_to_data", fake_image_to_data)
    evidence = read_document("scan-1", "scan.pdf", payload.getvalue())

    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == [
        "Shipper Averis\n77 Robinson",
        "Consignee Meridian",
    ]
    assert all(block.method == "ocr" for block in evidence.blocks)
    assert [location.page for location in evidence.blocks[0].locations] == [1, 1]
    assert [location.page for location in evidence.blocks[1].locations] == [1]
    assert evidence.blocks[0].locations[0].bbox is not None


def _minimal_text_pdf() -> bytes:
    stream = (
        b"BT /F1 12 Tf 72 720 Td (Shipping Instruction) Tj "
        b"0 -18 Td (Gross Weight: 123 kg) Tj ET"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode())
        payload.extend(obj)
        payload.extend(b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode())
    payload.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n"
        ).encode()
    )
    return bytes(payload)


def _descendant_worker(pipe: Connection, ignore_term: bool = False) -> None:
    _acquire_child_process_tree()
    child_code = (
        "import signal,time;"
        "signal.signal(signal.SIGTERM, signal.SIG_IGN);"
        "print('ready', flush=True);time.sleep(60)"
        if ignore_term
        else "import time;print('ready', flush=True);time.sleep(60)"
    )
    child = subprocess.Popen(
        [sys.executable, "-c", child_code],
        stdout=subprocess.PIPE,
        text=True,
        creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
    )
    assert child.stdout is not None
    assert child.stdout.readline() == "ready\n"
    pipe.send(child.pid)
    pipe.close()
    time.sleep(60)


def _process_exists(pid: int) -> bool:
    if sys.platform == "win32":
        result = subprocess.run(
            ["tasklist", "/FI", f"PID eq {pid}", "/FO", "CSV", "/NH"],
            check=False,
            capture_output=True,
            text=True,
            timeout=5,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
        return f'"{pid}"' in result.stdout
    process_stat = Path(f"/proc/{pid}/stat")
    if process_stat.exists():
        state = process_stat.read_text(encoding="utf-8").split()[2]
        return state != "Z"
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    return True


def test_process_tree_cleanup_terminates_a_spawned_descendant() -> None:
    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(target=_descendant_worker, args=(send,), daemon=True)
    process.start()
    send.close()
    descendant_pid: int | None = None
    try:
        assert receive.poll(10)
        descendant_pid = receive.recv()
    finally:
        receive.close()
        _stop_process_tree(process)

    assert descendant_pid is not None
    deadline = time.monotonic() + 5
    while _process_exists(descendant_pid) and time.monotonic() < deadline:
        time.sleep(0.05)
    assert not process.is_alive()
    assert not _process_exists(descendant_pid)


@pytest.mark.skipif(
    sys.platform == "win32",
    reason="POSIX process-group escalation has separate Windows Job Object coverage",
)
def test_process_tree_cleanup_kills_descendant_that_ignores_sigterm() -> None:
    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(
        target=_descendant_worker,
        args=(send, True),
        daemon=True,
    )
    process.start()
    send.close()
    descendant_pid: int | None = None
    try:
        assert receive.poll(10)
        descendant_pid = receive.recv()
    finally:
        receive.close()
        _stop_process_tree(process)

    assert descendant_pid is not None
    deadline = time.monotonic() + 5
    while _process_exists(descendant_pid) and time.monotonic() < deadline:
        time.sleep(0.05)
    assert not process.is_alive()
    assert not _process_exists(descendant_pid)


def test_docx_and_xlsx_multiline_label_values_normalize_equally() -> None:
    source_value = "Averis Trading\n77 Robinson Road"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Shipper (Principal or Seller) (发货人)"
    table.cell(0, 1).text = source_value
    docx_payload = BytesIO()
    document.save(docx_payload)

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Shipper", source_value])
    xlsx_payload = BytesIO()
    workbook.save(xlsx_payload)
    workbook.close()

    docx = read_document("docx", "bl.docx", docx_payload.getvalue())
    xlsx = read_document("xlsx", "si.xlsx", xlsx_payload.getvalue())

    assert docx.issues == []
    assert xlsx.issues == []
    assert normalize("shipper", docx.blocks[0].text) == normalize(
        "shipper", xlsx.blocks[0].text
    )
    assert normalize("shipper", docx.blocks[0].text) == (
        "averis trading 77 robinson road"
    )


def _formula_workbook(*, cache_is_current: bool) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = "Gross Weight (KG)"
    worksheet["B1"] = "=1-1"
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = not cache_is_current
    workbook.calculation.forceFullCalc = not cache_is_current
    raw = BytesIO()
    workbook.save(raw)
    workbook.close()

    source = ZipFile(BytesIO(raw.getvalue()))
    output = BytesIO()
    with source, ZipFile(output, "w", ZIP_DEFLATED) as destination:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                updated = data.replace(b"<f>1-1</f><v></v>", b"<f>1-1</f><v>0</v>")
                assert updated != data
                data = updated
            destination.writestr(member, data)
    return output.getvalue()


def test_current_cached_formula_zero_is_evidence() -> None:
    evidence = read_document(
        "current-formula",
        "current.xlsx",
        _formula_workbook(cache_is_current=True),
    )

    assert evidence.issues == []
    assert [block.text for block in evidence.blocks] == ["Gross Weight (KG): 0"]
    assert normalize("gross_weight_kg", evidence.blocks[0].text) == "0"


def test_stale_cached_formula_row_is_withheld_and_unresolved() -> None:
    evidence = read_document(
        "stale-formula",
        "stale.xlsx",
        _formula_workbook(cache_is_current=False),
    )

    assert evidence.blocks == []
    assert evidence.issues == [
        "xlsx_formula_values_uncertain:1:Sheet!B1",
        "document_has_no_readable_text",
    ]
