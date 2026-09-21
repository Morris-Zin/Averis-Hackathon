"""Adversarial acceptance checks for bounded document handling."""

from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from PIL import Image

from averis import documents
from averis.config import Settings
from averis.domain import AttachmentView, AuditEntry, CaseView, Classification
from averis.persistence import Base, Case, Database, Document, Run
from averis.processing import Processor
from averis.storage import Storage


def test_configured_document_limits_match_the_accepted_plan() -> None:
    assert documents.MAX_DOCUMENT_BYTES == 10 * 1024 * 1024
    assert documents.MAX_PDF_PAGES == 20
    assert documents.MAX_OCR_PAGES == 3
    assert documents.MAX_ARCHIVE_BYTES == 100 * 1024 * 1024
    assert documents.MAX_ARCHIVE_ENTRY_BYTES == 50 * 1024 * 1024
    assert documents.MAX_ARCHIVE_ENTRIES == 10_000
    assert documents.MAX_XLSX_SHEETS == 50
    assert documents.MAX_XLSX_ROWS_PER_SHEET == 20_000
    assert documents.MAX_XLSX_COLUMNS == 256
    assert documents.MAX_XLSX_CELLS == 200_000


def _blank_pdf(page_count: int) -> bytes:
    payload = BytesIO()
    images = [Image.new("RGB", (20, 20), "white") for _ in range(page_count)]
    try:
        images[0].save(
            payload,
            format="PDF",
            save_all=True,
            append_images=images[1:],
            resolution=72,
        )
    finally:
        for image in images:
            image.close()
    return payload.getvalue()


def _archive(entries: list[tuple[str, bytes]]) -> bytes:
    payload = BytesIO()
    with ZipFile(payload, "w", ZIP_DEFLATED) as archive:
        for name, content in entries:
            archive.writestr(name, content)
    return payload.getvalue()


def _workbook(cells: dict[str, object], *, second_sheet: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SI"
    for cell, value in cells.items():
        sheet[cell] = value
    if second_sheet:
        workbook.create_sheet("BL")["A1"] = "Bill of Lading"
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()
    return payload.getvalue()


def test_pdf_page_and_ocr_page_limits_are_explicit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[list[int]] = []

    def no_ocr(
        _document_id: str,
        _content: bytes,
        pages: list[int],
        _block_number: int,
    ) -> tuple[list[object], list[str]]:
        calls.append(pages)
        return [], []

    monkeypatch.setattr(documents, "_ocr_pdf_pages", no_ocr)

    page_limited = documents.read_document("pdf-pages", "scan.pdf", _blank_pdf(21))
    assert (
        f"pdf_page_limit_exceeded:21>{documents.MAX_PDF_PAGES}" in page_limited.issues
    )
    assert (
        f"pdf_ocr_page_limit_exceeded:{documents.MAX_PDF_PAGES}>"
        f"{documents.MAX_OCR_PAGES}"
    ) in page_limited.issues
    assert calls == [list(range(1, documents.MAX_OCR_PAGES + 1))]

    calls.clear()
    ocr_limited = documents.read_document("ocr-pages", "scan.pdf", _blank_pdf(4))
    assert (
        f"pdf_ocr_page_limit_exceeded:4>{documents.MAX_OCR_PAGES}" in ocr_limited.issues
    )
    assert calls == [[1, 2, 3]]


@pytest.mark.parametrize(
    ("constant", "limit", "entries", "expected"),
    [
        (
            "MAX_ARCHIVE_ENTRIES",
            1,
            [("a.xml", b"a"), ("b.xml", b"b")],
            "archive_entry_limit_exceeded:2>1",
        ),
        (
            "MAX_ARCHIVE_ENTRY_BYTES",
            4,
            [("large.xml", b"12345")],
            "archive_member_size_limit_exceeded:5>4",
        ),
        (
            "MAX_ARCHIVE_BYTES",
            5,
            [("a.xml", b"123"), ("b.xml", b"456")],
            "archive_size_limit_exceeded:6>5",
        ),
    ],
)
def test_office_archive_limits_are_explicit(
    monkeypatch: pytest.MonkeyPatch,
    constant: str,
    limit: int,
    entries: list[tuple[str, bytes]],
    expected: str,
) -> None:
    monkeypatch.setattr(documents, constant, limit)
    evidence = documents.read_document(
        "archive-limit", "shipment.docx", _archive(entries)
    )
    assert evidence.blocks == []
    assert evidence.issues == [expected]


def test_xlsx_structure_limits_are_explicit(monkeypatch: pytest.MonkeyPatch) -> None:
    with monkeypatch.context() as patch:
        patch.setattr(documents, "MAX_XLSX_SHEETS", 1)
        evidence = documents.read_document(
            "sheet-limit",
            "shipment.xlsx",
            _workbook({"A1": "Shipping Instruction"}, second_sheet=True),
        )
        assert "xlsx_sheet_limit_exceeded:2>1" in evidence.issues

    with monkeypatch.context() as patch:
        patch.setattr(documents, "MAX_XLSX_ROWS_PER_SHEET", 2)
        evidence = documents.read_document(
            "row-limit", "shipment.xlsx", _workbook({"A3": "Consignee: Example"})
        )
        assert "xlsx_row_limit_exceeded:SI:3" in evidence.issues

    with monkeypatch.context() as patch:
        patch.setattr(documents, "MAX_XLSX_COLUMNS", 2)
        evidence = documents.read_document(
            "column-limit", "shipment.xlsx", _workbook({"C1": "Port: Klang"})
        )
        assert "xlsx_column_limit_exceeded:SI:3" in evidence.issues

    with monkeypatch.context() as patch:
        patch.setattr(documents, "MAX_XLSX_CELLS", 4)
        evidence = documents.read_document(
            "cell-limit",
            "shipment.xlsx",
            _workbook({"A1": "SI", "B3": "Gross weight: 100 kg"}),
        )
        assert "xlsx_cell_scan_limit_exceeded:SI:6>4" in evidence.issues


def test_reader_and_preview_deadlines_fail_explicitly() -> None:
    evidence = documents.read_document_bounded(
        "deadline", "shipment.txt", b"Shipper: Example", timeout_seconds=1e-9
    )
    assert evidence.blocks == []
    assert evidence.issues == ["document_read_timeout:1e-09s"]

    with pytest.raises(
        documents.DocumentPreviewError, match=r"^preview_timeout:1e-09s$"
    ):
        documents.render_preview_bounded(
            "shipment.pdf", _blank_pdf(1), timeout_seconds=1e-9
        )


def test_preview_rejects_bad_pages_timeouts_and_corrupt_pdf() -> None:
    payload = _blank_pdf(1)
    with pytest.raises(
        documents.DocumentPreviewError, match="preview_page_out_of_range:0"
    ):
        documents.render_preview_bounded("shipment.pdf", payload, page=0)
    with pytest.raises(documents.DocumentPreviewError, match="preview_invalid_timeout"):
        documents.render_preview_bounded("shipment.pdf", payload, timeout_seconds=0)
    with pytest.raises(documents.DocumentPreviewError, match="preview_unreadable"):
        documents.render_preview_bounded("shipment.pdf", b"not a PDF")


def test_image_reader_and_preview_reject_format_mismatch_and_pixel_excess() -> None:
    png = Image.new("RGB", (20, 20), "white")
    png_payload = BytesIO()
    png.save(png_payload, format="PNG")
    png.close()

    mislabeled = documents.read_document(
        "mislabeled", "scan.jpg", png_payload.getvalue()
    )
    assert mislabeled.blocks == []
    assert mislabeled.issues == ["document_unreadable:invalid_jpg"]
    with pytest.raises(
        documents.DocumentPreviewError,
        match="preview_unreadable:invalid_jpg",
    ):
        documents.render_preview_bounded("scan.jpg", png_payload.getvalue())

    oversized = Image.new("RGB", (documents.MAX_RENDER_DIMENSION + 1, 1), "white")
    oversized_payload = BytesIO()
    oversized.save(oversized_payload, format="PNG")
    oversized.close()
    evidence = documents.read_document(
        "oversized-image", "scan.png", oversized_payload.getvalue()
    )
    expected = (
        f"image_pixel_limit_exceeded:{documents.MAX_RENDER_DIMENSION + 1}x1>"
        f"{documents.MAX_RENDER_PIXELS}"
    )
    assert evidence.blocks == []
    assert evidence.issues == [expected]
    with pytest.raises(
        documents.DocumentPreviewError,
        match=(
            f"preview_image_pixel_limit_exceeded:"
            f"{documents.MAX_RENDER_DIMENSION + 1}x1>"
            f"{documents.MAX_RENDER_PIXELS}"
        ),
    ):
        documents.render_preview_bounded("scan.png", oversized_payload.getvalue())


class _ComparisonWithoutExtraction:
    def __init__(self) -> None:
        self.extract_calls = 0

    def classify(
        self,
        _subject: str,
        _body: str,
        *,
        attachment_filenames: tuple[str, ...] = (),
        load_attachment_previews=None,
    ) -> Classification:
        return Classification(
            suggested="BL_COMPARISON",
            accepted="BL_COMPARISON",
            confidence=1,
            probabilities={"BL_COMPARISON": 1},
            source="fixture",
            model="limit-test",
        )

    def extract(self, _evidence: object) -> object:
        self.extract_calls += 1
        raise AssertionError("unsupported input must not be sent for extraction")


def test_unsupported_document_completes_with_visible_review_issue(tmp_path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'limits.db'}",
        storage_backend="local",
        storage_dir=str(tmp_path / "objects"),
        tasks_queue="",
        live_enabled=True,
        budget_verified=True,
    )
    database = Database(settings.database_url)
    Base.metadata.create_all(database.engine)
    storage = Storage(settings)
    object_key, digest = storage.put(b"unsupported source")
    case_id, run_id, document_id = "case-limit", "run-limit", "document-limit"
    view = CaseView(
        id=case_id,
        subject="Review unsupported document",
        sender="fixture@example.test",
        body="Compare the attachment.",
        received_at="2026-09-20T00:00:00+00:00",
        revision=1,
        input_revision=1,
        processing="queued",
        stage="queued",
        workflow="open",
        assignee="Unassigned",
        review_reasons=[],
        attachments=[AttachmentView(id=document_id, filename="shipment.rtf")],
        history=[
            AuditEntry(
                at="2026-09-20T00:00:00+00:00",
                actor="fixture",
                action="created",
                detail="unsupported document acceptance fixture",
            )
        ],
    )
    with database.session() as session, session.begin():
        session.add(
            Case(
                id=case_id,
                workspace_id="workspace-limit",
                revision=1,
                input_revision=1,
                active_run_id=run_id,
                state=view.model_dump(mode="json"),
            )
        )
        session.add(
            Document(
                id=document_id,
                workspace_id="workspace-limit",
                case_id=case_id,
                filename="shipment.rtf",
                object_key=object_key,
                sha256=digest,
                size=len(b"unsupported source"),
            )
        )
        session.add(
            Run(
                id=run_id,
                case_id=case_id,
                input_revision=1,
                purpose="development",
                status="queued",
            )
        )

    intelligence = _ComparisonWithoutExtraction()
    processor = Processor(database, settings, storage, factory=lambda *_: intelligence)
    try:
        assert processor.execute(run_id) == "completed"
        with database.session() as session:
            row = session.get(Case, case_id)
            run = session.get(Run, run_id)
            assert row is not None
            assert run is not None
            saved = CaseView.model_validate(row.state)
            assert run.status == "completed"
            assert saved.processing == "completed"
            assert saved.report is None
            assert saved.attachments[0].evidence is not None
            assert saved.attachments[0].evidence.issues == [
                "unsupported_document_type:.rtf"
            ]
            assert "unsupported_document_type:.rtf" in saved.review_reasons
            assert row.needs_review
            assert intelligence.extract_calls == 0
    finally:
        database.engine.dispose()
